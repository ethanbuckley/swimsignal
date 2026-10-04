"""Write the two small workbooks test_sw_history.py reads: rows cut, unchanged, from Scottish Water's
overflow event workbooks of 5 and 31 August 2026 (sw_history.FILES), with their sheet names and
the two header rows of each sheet. Date cells are kept as date cells, text dates as text.

Contains Scottish Water data licensed under the Open Government Licence v3.0.

The rows were picked to cover: an overflow listed with "No Events" and "No Data"; Ellon WWTW CSO,
published as STW001545 up to 2023 and CSO006851 from 2024; Erskine's two events inside the missing
hour of 31 March 2024; Helensburgh's event from the repeated hour of 29 October 2023; text dates
(reported 2025, non-reported 2024); Fort William's 2022 day records without a duration; a 15-minute
event; Lincluden, STW000468 in 2022 and CSO007138 in the summary.

    uv run python tests/fixtures/scottish_water/make_fixture.py
"""

from datetime import datetime, timedelta
from pathlib import Path

from openpyxl import Workbook

HERE = Path(__file__).parent
EPOCH = datetime(1899, 12, 30)  # noqa: DTZ001  Excel's day 0
N = None
REP_SUM_TOP = [N] * 14 + ["No. of Days Data", N, N, "Total Number of Overflow Events", N, N, N, N,
                          "Total Duration of Overflow Events\r\n([h]:mm:ss)", N, N, N, N,
                          "Total Volume of Overflow Events\r\n(m3)", N, N, N, N, N]
REP_SUM_HEAD = ["Permit Number", "Asset ID", "Asset Name", "SEPA Measurement Point Description",
                "SW Unique Measurement Point Description", "Scottish Water Area", "Network or WWTW",
                "Type of Overflow CSO/SSSO/EO", "Bathing Season Only Reporting?", "Asset Postcode",
                "Local Authority Area", "X Co-ordinate", "Y Co-ordinate", "Discharge Receiving Water",
                2025, 2024, 2023, 2025, 2024, 2023, 2022, 2021, 2025, 2024, 2023, 2022, 2021,
                2025, 2024, 2023, 2022, 2021, "Comments"]
NR = "Not Required"
REP_EV_HEAD = ["Permit Number", "Ellipse No.", "Asset ID", "Asset Name", "SEPA Measurement Point Description",
               "SW Measurement Point Description", "Area", "Network or WWTW", "No. of Days Data",
               "Start of Overflow Event", "End of Overflow Event", "Duration of Overflow Event\r\n[h]:mm:ss",
               "Volume Discharged \r\n(m3)", "Comments"]
LIC = ["Licence Number"] + REP_EV_HEAD[1:]
REP_2022_HEAD = LIC[:8] + LIC[9:]
REP_2021_HEAD = ["Licence Number", "Asset ID", "Asset Name", "SEPA Measurement Point Description",
                 "SW Unique Measurement Point Description", "Area", "Network or WWTW",
                 "Start of Overflow Event", "End of Overflow Event", "Duration of Overflow Event\r\n[h]:mm:ss",
                 "Volume Discharged \r\n(m3)", "Comments"]
AIRDRIE = ["CAR/L/1026106", 5000594132, "CSO005882", "67 SOUTH COMMONHEAD CSO 2009 NS766664", "CSO event"]
ELLON25 = ["CAR/L/1003861", 5001010360, "CSO006851", "ELLON WWTW CSO AND EO", "CSO event", "Ellon WWTW CSO Event",
           "East", "WWTW", 365]
ELLON24 = ["CAR/L/1003861", 5001010360, "CSO006851", "ELLON WWTW CSO NJ966305", "CSO event", "Ellon WWTW CSO Event",
           "East", "WWTW", 366]
ELLON23 = ["CAR/L/1003861", 5000062284, "STW001545", "ELLON WWTW 1989 NJ966305", "CSO event", "Ellon WWTW CSO event",
           "East", "WWTW", 339]
IRON = ["CAR/L/1001146", 5001012870, "CSO006855", "IRON MILL BAY WWTW INLET CSO NT061843",
        "Ironmill Bay WWTW CSO storm event", "Ironmill Bay WWTW CSO Event", "East", "WWTW", 365]
ERSK = ["CAR/L/1008817", 5001065151, "CSO007524", "ERSKINE WWTW - STORM TANK SSSO",
        "Erskine WWTW SSSO (3DWF) storm event", "Erskine WWTW SSSO Event", "West", "WWTW", 366]
FORT = ["CAR/L/1001897", 5000062377, "STW001638", "FORT WILLIAM WWTW NN108750", "Fort William WWTW CSO storm event",
        "Fort William WWTW CSO event", "North", "WWTW"]
ELLON21 = ["CAR/L/1003861", "STW001545", "ELLON WWTW 1989 NJ966305", "CSO event", "Ellon WWTW CSO event", "East", "WWTW"]

REPORTED = {
    "Overflow Summary 2021-2025": [REP_SUM_TOP, REP_SUM_HEAD,
        ["CAR/L/1026106", "CSO005882", "67 SOUTH COMMONHEAD CSO 2009 NS766664", "CSO event",
         "Airdrie, 67 South Commonhead CSO Event", "South", "Network", "CSO", N, "ML6 6QR", "North Lanarkshire",
         276564, 666431, "North Burn (culverted)", 206, 0, 0, "No Events", "No Data", "No Data", "No Events",
         "No Events", "No Events", "No Data", "No Data", "No Events", "No Events", NR, NR, NR, NR, NR, N],
        ["CAR/L/1003861", "CSO006851", "ELLON WWTW CSO AND EO", "CSO event", "Ellon WWTW CSO Event", "East", "WWTW",
         "CSO", N, "AB41 9EY", "Aberdeenshire", 396640, 830488, "River Ythan", 365, 366, 339, 2, 5, 5, 9, 2,
         0.07188657407407407, 2.765011574074074, 2.9479166666666665, 1.208460648148148, 1.115173611111111,
         NR, NR, NR, NR, NR, N],
        ["CAR/L/1008817", "CSO007524", "ERSKINE WWTW - STORM TANK SSSO", "Erskine WWTW SSSO (3DWF) storm event",
         "Erskine WWTW SSSO Event", "West", "WWTW", "SSSO", N, "PA4 9PB", "Renfrewshire", 249358, 669099,
         "Clyde Estuary", 365, 366, 365, 312, 258, 160, 116, 94, 95.5092592592593, 133.89092592592598,
         112.82057870370369, 109.52083333333333, 81.02083333333334, "No Data", "No Data", "No Data", "No Data",
         "No Data", N],
        ["CAR/L/1003465", "CSO007529", "HELENSBURGH WWTW - POST-PST SSSO", "SSSO event", "Helensburgh WWTW SSSO Event",
         "West", "WWTW", "SSSO", N, "G82 5HG", "Argyll and Bute", 232054, 679869, "Clyde Estuary", 365, 366, 365,
         1697, 1722, 1521, 616, 553, 288.8301851851855, 327.8403935185187, 320.2521990740737, 89.96875,
         91.10416666669335, 2761762.0100000035, 490427.60079654166, 944484.3021650796, 1302710.6633606304,
         1274477.7679909095, N],
        ["CAR/L/1001146", "CSO006855", "IRON MILL BAY WWTW INLET CSO NT061843", "Ironmill Bay WWTW CSO storm event",
         "Ironmill Bay WWTW CSO Event", "East", "WWTW", "CSO", N, "KY11 3DP", "Fife", 306064, 684294,
         "Forth Estuary", 365, 366, 364, 3, 5, 9, 9, 33, 0.29337962962962966, 1.3726273148148147,
         1.4767476851851853, 1.2300578703703704, 5.1996064814814815, 3460.3156735374077, 13629.583050787107,
         5136.96, 6853.68, 24699.599999999988, N]],
    "Overflow Events in 2025": [["Year", 2025] + [N] * 12, REP_EV_HEAD,
        AIRDRIE + ["Airdrie, 67 South Commonhead CSO Event", "South", "Network", 206, "No Events", "No Events",
                   "No Events", NR, "Intermittent data gaps throughout the year."],
        ELLON25 + [45815.5390625, 45815.553391203706, 0.014328703703703703, NR, N],
        ELLON25 + [45998.916666666664, 45998.974224537036, 0.05755787037037037, NR, N],
        IRON + ["25/05/2025 02:06:15", "25/05/2025 03:28:34", 0.05716435185604496, 1628.3241781463585, N],
        IRON + ["26/05/2025 21:56:08", "27/05/2025 02:09:41", 0.17607638888875954, 1649.5406681526558, N],
        IRON + ["03/10/2025 15:52:38", "03/10/2025 17:19:14", 0.06013888888992369, 182.4508272383939, N]],
    "Overflow Events in 2024": [["Year", 2024] + [N] * 12, LIC,
        AIRDRIE + ["Airdrie, 67 South Commonhead CSO Event", "South", "Network", 0, "No Data", "No Data", "No Data",
                   NR, "No data available for 2024. Monitoring improvement project ongoing."],
        ELLON24 + [45293.83888888889, 45293.879166666666, 0.04010416666666667, NR, N],
        ELLON24 + [45294.08541666667, 45295.785416666666, 1.7, NR, N],
        ERSK + [45382.055555555555, 45382.06805555556, 0.012499999999999999, N, N],
        ERSK + [45382.080555555556, 45382.092361111114, 0.011574074074074075, N, N]],
    "Overflow Events in 2023": [["Year", 2023] + [N] * 12, LIC,
        AIRDRIE + ["Airdrie, 67 South Commonhead CSO event", "South", "Network", 0, "No Data", "No Data", "No Data",
                   NR, "No data available for 2023.  Monitoring improvement project ongoing."],
        ELLON23 + [45054.99444444444, 45055.01666666667, 0.022569444444444444, NR, N],
        ELLON23 + [45267.884722222225, 45268.760416666664, 0.8750810185185185, NR, N],
        ELLON23 + [45283.59930555556, 45283.73333333333, 0.13402777777777777, NR, N],
        ELLON23 + [45287.48541666667, 45288.65694444445, 1.1715277777777777, NR, N],
        ELLON23 + [45290.76527777778, 45291.510416666664, 0.7447106481481481, NR, N],
        ["CAR/L/1003465", 5000004318, "STW000372", "HELENSBURGH WWTW 2001 NS321800", "SSSO event",
         "Helensburgh WWTW SSSO event", "West", "WWTW", 365, 45228.041666666664, 45229.53194444445,
         1.4900925925925925, 427.696963, N]],
    "Overflow Events in 2022": [["Year", 2022] + [N] * 11, REP_2022_HEAD,
        AIRDRIE + ["Airdrie, 67 South Commonhead CSO event", "South", "Network", "No Events", "No Events",
                   "No Events", NR, N],
        FORT + [44562, 44562, "No Data Available", 453, ("Asset returned to Scottish Water during 2022.\r\n\r\n"
                "Duration of events is not available.\r\n\r\nResubmission 12/11/25 – Due to current monitor "
                "limitations overflow start/end times are unable to be provided.  Therefore, start/end times have "
                "been removed.")],
        FORT + [44563, 44563, "No Data Available", 7224, N]],
    "Overflow Events in 2021": [["Year", 2021] + [N] * 10, REP_2021_HEAD,
        ["CAR/L/1026106", "CSO005882", "67 SOUTH COMMONHEAD CSO 2009 NS766664", "CSO event",
         "Airdrie, 67 South Commonhead CSO event", "South", "Network", "No events", "No events", "No events",
         "Not required", N],
        ELLON21 + [44264.49166666667, 44265.575694444444, 1.0839236111111112, "Not required", N],
        ELLON21 + [44534.635416666664, 44534.666666666664, 0.03125, "Not required", N]],
}

NON_SUM_TOP = [N] * 13 + ["No. of Days Data", N, N, N, "Total Number of Overflow Events", N, N, N,
                          "Total Duration of Overflow Events\r\n([h]:mm:ss)", N, N, N, N]
NON_SUM_HEAD = ["Permit Number", "Asset ID", "Asset Name", "SW Unique Measurement Point Description",
                "Scottish Water Area", "Network or WWTW", "Type of Overflow CSO/SSSO/EO", "Postcode",
                "Local Authority Area", "X Co-ordinate", "Y Co-ordinate", "Discharge Receiving Water",
                "All Events Published?", 2025, 2024, 2023, 2022, 2025, 2024, 2023, 2022, 2025, 2024, 2023, 2022,
                "Comments"]
NON_EV_HEAD = ["Permit Number", "Ellipse No.", "Asset ID", "Asset Name", "SW Unique Measurement Point Description",
               "Area", "Network or WWTW", "No. of Days Data", "Start of Overflow Event", "End of Overflow Event",
               "Duration of Overflow Event\r\n[h]:mm:ss", "Comments"]
NON_OLD_HEAD = ["Licence Number", "Ellipse No.", "Asset ID", "Asset Name", "SW Measurement Point Description"] + NON_EV_HEAD[5:]
ABERD = ["CAR/L/1026105", 5000244962, "CSO004465", "ABERDEEN, 143 GEORGE ST NJ93069510 CSO",
         "Aberdeen, 143 George St CSO Event", "East", "Network", 365]
NO_DATA = ["No Data"] * 3

NON_REPORTED = {
    "Overflow Summary 2022-2025": [NON_SUM_TOP, NON_SUM_HEAD,
        ["CAR/L/1004222", "CSO007267", "ABERCHIRDER WWTW SSSO", "Aberchirder WWTW SSSO Event", "East", "WWTW", "SSSO",
         "AB54 7SS", "Aberdeenshire", 362919, 852078, "Arkland Burn", N, 0, 0, 0, 0] + ["No Data"] * 8 + [N],
        ["CAR/L/1026105", "CSO004465", "ABERDEEN, 143 GEORGE ST NJ93069510 CSO", "Aberdeen, 143 George St CSO Event",
         "East", "Network", "CSO", "AB25 1BQ", "Aberdeen City", 393960, 806586, "Un-named Tributary of Den Burn",
         "Yes\r\nFrom 2025", 365, 0, 0, 0, 4, "No Data", "No Data", "No Data", 0.05555555555555555, "No Data",
         "No Data", "No Data", "2025 Data Published August 2026."],
        ["CAR/L/1026153", "CSO001640", "STONEYBURN, BENTS FARM/CANNOP CRES CSO",
         "Bathgate, Stoneyburn Bents Farm/Cannop Cres CSO Event", "South", "Network", "CSO", "EH55 8JF",
         "West Lothian", 296862.227, 662079.403, "Breich Water", "No\r\nEvents >15:30 Only", 365, 291, 356, 306,
         2, 1, 0, 0, 0.2881944444444444, 0.017361111111111112, 0, 0, N],
        ["CAR/L/1003532", "CSO007138", "LINCLUDEN WWTW SSSO", "Lincluden WWTW SSSO Event", "South", "WWTW", "SSSO",
         "DG2 0DL", "Dumfries and Galloway", 296306, 578026, "Cluden Water", "No\r\nEvents >15:30 Only", 204, 366,
         365, 365, 370, 440, 755, 444, 111.9411574074073, 130.4345254629629, 247.29151620370368,
         216.04204861111103, N]],
    "Overflow Events in 2025": [["Year", 2025] + [N] * 10, NON_EV_HEAD,
        ["CAR/L/1004222", 5001047762, "CSO007267", "ABERCHIRDER WWTW SSSO", "Aberchirder WWTW SSSO Event", "East",
         "WWTW", 0] + NO_DATA + [N],
        ABERD + [45815.51388888889, 45815.524305555555, 0.010416666664241347, N],
        ABERD + [45822.5, 45822.524305555555, 0.024305555554747116, ""],
        ABERD + [45830.05902777778, 45830.069444444445, 0.010416666664241347, ""],
        ABERD + [45865.600694444445, 45865.61111111111, 0.010416666664241347, ""]],
    "Overflow Events in 2024": [["Year", 2024] + [N] * 10, ["Licence Number"] + NON_EV_HEAD[1:],
        ["CAR/L/1004222", 5001047762, "CSO007267", "ABERCHIRDER WWTW SSSO", "Aberchirder WWTW SSSO Event", "EAST",
         "WWTW", 0] + NO_DATA + ["Monitoring improvement project ongoing. "],
        ["CAR/L/1026153", 5000229444, "CSO001640", "STONEYBURN, BENTS FARM/CANNOP CRES CSO",
         "Bathgate, Stoneyburn Bents Farm/Cannop Cres CSO Event", "South", "Network", 291, "31/12/2024 09:40:00",
         "31/12/2024 10:05:00", 0.01736111110949423, "New monitor installed.  Data available from 16/03/24."]],
    "Overflow Events in 2023": [["Year", 2023] + [N] * 10, NON_OLD_HEAD],
    "Overflow Events in 2022": [["Year", 2022] + [N] * 10, NON_OLD_HEAD,
        ["CAR/L/1003532", 5000005230, "STW000468", "LINCLUDEN WWTW 1930 NX963780", "LINCLUDEN WWTW SSSO event",
         "South", "WWTW", 365, 44864.078310185185, 44864.26857638889, 0.1902662037027767,
         "Based on 365 days where data is available."]],
}


def write(sheets: dict, path: Path) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        head = rows[1]
        when = {i for i, h in enumerate(head) if str(h).startswith(("Start of", "End of"))}
        dur = {i for i, h in enumerate(head) if str(h).startswith("Duration")}
        for r in rows:
            ws.append(r)
        for row in ws.iter_rows(min_row=3):
            for c in row:
                if c.column - 1 in when and isinstance(c.value, int | float):
                    c.value = EPOCH + timedelta(days=c.value)
                    c.number_format = "dd/mm/yyyy hh:mm:ss"
                elif c.column - 1 in dur and isinstance(c.value, int | float):
                    c.number_format = "[h]:mm:ss"
    wb.save(path)


if __name__ == "__main__":
    write(REPORTED, HERE / "reported.xlsx")
    write(NON_REPORTED, HERE / "non_reported.xlsx")
