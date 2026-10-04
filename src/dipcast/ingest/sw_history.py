"""Scottish Water's published overflow event history, one row per event (Scotland plan task S1).

Two spreadsheets from Scottish Water's document hub, both under the Open Government Licence v3.0:

- reported to SEPA, 2021-2025: the overflows with annual reporting in their permits (about 300
  measurement points); events longer than 15 min 30 s only.
- not reported to SEPA, 2022-2025: other monitored overflows, where Scottish Water verified the
  data; all events, short ones too, from 2025 for the overflows added in August 2026.

Contains Scottish Water data licensed under the Open Government Licence v3.0.

Each workbook has a summary sheet (one row per measurement point, with the current Asset ID and
a British National Grid position) and one "Overflow Events in <year>" sheet per year. What the
files leave open, and what this module does about it:

- Time zone. Times are "dd/mm/yyyy hh:mm:ss" with no zone given. Starts are read as Europe/London
  clock time (inferred, 4 Oct 2026): across both workbooks and 2021-2025, the missing hour of the
  March clock change (01:00-02:00) holds 15 starts or ends against 30 and 23 in the two hours
  after it, and the repeated hour of October holds 56 against 40 and 33. Neither hour is empty
  or doubled, so some monitors may log GMT. A start inside the missing hour is read as GMT; one
  in the repeated hour as BST. The end is the start plus the clock-time difference, because the
  published durations are that difference, across clock changes too (checked).
- Asset IDs. Up to 2023 many overflows at treatment works carry the works' STW... id; from 2024
  they carry their own CSO... id, which the live API uses. Each row takes the Asset ID of the
  summary row with the same measurement point description (own workbook first, then the other),
  after case, punctuation and the word "event" are dropped; else its own. The published id is
  kept in `asset_id_published`.
- Day records. Some reported overflows have rows whose start equals their end, at midnight:
  a day with a discharge, and in most a duration, but no time of day. They keep the start
  (midnight local), no end, and the published duration.
- Text cells. Some sheets hold dates as text instead of date cells; both are read. Rows saying
  "No Events", "No Data" or "Not Required" become statuses, not events.
- Like edm_events, an event of 60 days or more is dropped (12 rows in the August 2026 files),
  and so is an exact repeat of a row (1).

`site_name` is the measurement point description. `events()` returns the `edm_events` columns plus `asset_id_published`, `year` (the sheet's) and
`duration_published_h`; `yearly()` adds one row per overflow and year with the 12/24-hour count.
"""

from __future__ import annotations

import logging
import re
from datetime import datetime, time, timedelta
from pathlib import Path

import numpy as np
import pandas as pd
from pyproj import Transformer

from dipcast import config
from dipcast.ingest.edm_events import COLS

log = logging.getLogger(__name__)

RAW_DIR = config.RAW / "scottish_water"
PAGE = "https://www.scottishwater.co.uk/Your-Home/Your-Waste-Water/Overflows/Overflow-Event-Data"
BASE = ("https://www.scottishwater.co.uk/-/media/scottishwater/document-hub/key-publications/"
        "improving-urban-waters/cso-data/")
# source -> (file name, size in bytes on 4 Oct 2026). A different size means a new version.
FILES = {
    "Scottish Water reported 2021-2025":
        ("sw-reported-overflow-event-data-to-sepa-20212025--summary--v2-050826.xlsx", 5_454_175),
    "Scottish Water non-reported 2022-2025":
        ("scottish-water-non-reported-overflow-event-data-2022-2025-and-summary--310826.xlsx", 7_549_740),
}
CREDIT = "Contains Scottish Water data licensed under the Open Government Licence v3.0."
TZ = "Europe/London"
SHORT_S = 15 * 60 + 30          # the reported workbook keeps only events longer than this
STATUSES = ["events", "no events", "no data", "not required"]   # best first, per overflow-year
EXTRA = ["asset_id_published", "year", "duration_published_h"]

_EXCEL_EPOCH = datetime(1899, 12, 30)   # noqa: DTZ001  Excel's day 0; cells hold clock time
_TEXT_TIME = re.compile(r"^(\d{1,2})/(\d{1,2})/(\d{4})\s+(\d{1,2}):(\d{2})(?::(\d{2}))?$")
_TEXT_DURATION = re.compile(r"^(\d+):(\d{2}):(\d{2})$")
_TO_WGS = Transformer.from_crs(27700, 4326, always_xy=True)


def clock_time(v) -> pd.Timestamp:
    """A start or end cell as naive clock time: a date cell, an Excel serial number, or text
    "dd/mm/yyyy hh:mm[:ss]". Anything else ("No Events", "No Data", blank) is NaT."""
    if isinstance(v, datetime):
        return pd.Timestamp(v).round("s")
    if isinstance(v, int | float) and not isinstance(v, bool) and v > 0:
        return pd.Timestamp(_EXCEL_EPOCH + timedelta(days=float(v))).round("s")
    if isinstance(v, str) and (m := _TEXT_TIME.match(v.strip())):
        d, mo, y, h, mi, s = (int(x or 0) for x in m.groups())
        return pd.Timestamp(y, mo, d, h, mi, s)
    return pd.NaT


def duration_s(v) -> float:
    """A "[h]:mm:ss" duration cell in seconds; NaN when it holds text such as "No Data Available"."""
    if isinstance(v, timedelta):
        return float(round(v.total_seconds()))
    if isinstance(v, datetime):                     # a duration of a day or more read as a date
        return float(round((v - _EXCEL_EPOCH).total_seconds()))
    if isinstance(v, time):
        return float(round(v.hour * 3600 + v.minute * 60 + v.second + v.microsecond / 1e6))
    if isinstance(v, int | float) and not isinstance(v, bool):
        return float(round(v * 86400))
    if isinstance(v, str) and (m := _TEXT_DURATION.match(v.strip())):
        h, mi, s = (int(x) for x in m.groups())
        return float(h * 3600 + mi * 60 + s)
    return np.nan


def point_key(desc) -> str:
    """Measurement point description without case, punctuation or the word "event": the same
    point is spelt "Ellon WWTW CSO event" in 2023 and "Ellon WWTW CSO Event" in 2025, and
    "Greenock, No.5" in one year and "Greenock No5" in another."""
    s = re.sub(r"[.']", "", str(desc or "").lower().replace("&", " and "))
    s = re.sub(r"[^a-z0-9]+", " ", s)
    return " ".join(w for w in s.split() if w not in ("event", "events"))


def status(v) -> str:
    t = str(v or "").strip().lower()
    if t in ("no events", "no spills"):
        return "no events"
    if t.startswith("no data"):
        return "no data"
    return "not required"                           # "Not Required", "Not Reported", blank


def _sheet(ws) -> tuple[list[str], list[tuple]]:
    """Header (second row) and the non-empty rows below it. The first row is a title ("Year",
    2025) on event sheets and group headings on the summary."""
    rows = ws.iter_rows(values_only=True)
    next(rows, None)
    header = [str(h).strip() if h is not None else "" for h in next(rows, ())]
    return header, [r for r in rows if any(v is not None and v != "" for v in r)]


def _col(header: list[str], *starts: str) -> int | None:
    for i, h in enumerate(header):
        if any(h.lower().startswith(s.lower()) for s in starts):
            return i
    return None


_POINT = ("SW Unique Measurement Point", "SW Measurement Point")


def read_workbook(path: Path) -> tuple[pd.DataFrame, pd.DataFrame]:
    """(summary, rows) of one workbook. `rows` holds every row of every event sheet: events and
    the "No Events"/"No Data" lines, with naive clock times."""
    from openpyxl import load_workbook  # dev dependency; only this off-CI path needs it

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        summary, rows = None, []
        for name in wb.sheetnames:
            if name.startswith("Overflow Summary"):
                summary = _summary(*_sheet(wb[name]))
            elif m := re.match(r"Overflow Events in (\d{4})$", name):
                rows.append(_events_sheet(int(m.group(1)), *_sheet(wb[name])))
    finally:
        wb.close()
    if summary is None or not rows:
        raise ValueError(f"{path.name}: no summary or event sheets")
    return summary, pd.concat(rows, ignore_index=True)


def _summary(header: list[str], rows: list[tuple]) -> pd.DataFrame:
    ix = {k: _col(header, *s) for k, s in {
        "asset_id": ("Asset ID",), "point": _POINT, "overflow_type": ("Type of Overflow",),
        "local_authority": ("Local Authority",), "receiving_water": ("Discharge Receiving Water",),
        "x": ("X Co-ordinate",), "y": ("Y Co-ordinate",)}.items()}
    df = pd.DataFrame({k: [r[i] if i is not None and i < len(r) else None for r in rows] for k, i in ix.items()})
    df = df[df["asset_id"].notna()].reset_index(drop=True)
    df["key"] = df["point"].map(point_key)
    x, y = pd.to_numeric(df["x"], errors="coerce"), pd.to_numeric(df["y"], errors="coerce")
    lon, lat = _TO_WGS.transform(x.to_numpy(float), y.to_numpy(float))
    df["lat"], df["lon"] = np.round(lat, 6), np.round(lon, 6)
    return df


def _events_sheet(year: int, header: list[str], rows: list[tuple]) -> pd.DataFrame:
    ia, ip = _col(header, "Asset ID"), _col(header, *_POINT)
    i_s, i_e = _col(header, "Start of Overflow Event"), _col(header, "End of Overflow Event")
    i_d, i_days = _col(header, "Duration of Overflow Event"), _col(header, "No. of Days Data")
    if None in (ia, ip, i_s, i_e, i_d):
        raise ValueError(f"Overflow Events in {year}: columns not found in {header}")
    cell = lambda r, i: r[i] if i is not None and i < len(r) else None
    df = pd.DataFrame({
        "year": year,
        "asset_id_published": [cell(r, ia) for r in rows],
        "point": [cell(r, ip) for r in rows],
        "start": pd.to_datetime([clock_time(cell(r, i_s)) for r in rows]),
        "end": pd.to_datetime([clock_time(cell(r, i_e)) for r in rows]),
        "duration_pub_s": [duration_s(cell(r, i_d)) for r in rows],
        "days_data": pd.to_numeric(pd.Series([cell(r, i_days) for r in rows], dtype="object"), errors="coerce"),
        "status": ["events" if not pd.isna(clock_time(cell(r, i_s))) else status(cell(r, i_s)) for r in rows],
    })
    df = df[df["asset_id_published"].notna()]
    df["key"] = df["point"].map(point_key)
    return df


def london_to_utc(naive: pd.Series, dst: bool = True) -> pd.Series:
    """Europe/London clock time to UTC. In October's repeated hour `dst` picks BST (True) or GMT;
    a time inside March's missing hour is read as GMT (moved on one hour, then converted)."""
    local = naive.dt.tz_localize(TZ, ambiguous=np.full(len(naive), dst), nonexistent=pd.Timedelta(hours=1))
    return local.dt.tz_convert("UTC")


def assign_sites(rows: pd.DataFrame, summaries: list[pd.DataFrame]) -> pd.DataFrame:
    """Current Asset ID, name and position per row, from the first summary (own workbook first)
    with the same measurement point; else the summary row with the same Asset ID; else the
    published id and no position."""
    rows = rows.copy()
    sid = pd.Series(pd.NA, index=rows.index, dtype="object")
    for s in summaries:
        sid = sid.fillna(rows["key"].map(s.drop_duplicates("key").set_index("key")["asset_id"]))
    rows["site_id"] = sid.fillna(rows["asset_id_published"]).astype(str)
    by_key = pd.concat(summaries).drop_duplicates("key").set_index("key")
    by_id = pd.concat(summaries).drop_duplicates("asset_id").set_index("asset_id")
    for c in ("lat", "lon"):
        rows[c] = rows["key"].map(by_key[c]).fillna(rows["site_id"].map(by_id[c]))
    return rows


def events(rows: pd.DataFrame, source: str) -> pd.DataFrame:
    """One row per published event, `edm_events` columns first. `rows` comes from assign_sites."""
    ev = rows[rows["status"] == "events"].copy()
    n = len(ev)
    ev = ev.drop_duplicates(["site_id", "key", "start", "end"])
    if len(ev) < n:
        log.info("%s: dropped %d repeated rows", source, n - len(ev))
    day = ev["start"] == ev["end"]                  # a day with a discharge, no time of day
    ev["event_start"] = london_to_utc(ev["start"])
    # The end keeps its clock-time distance from the start: the published durations are that
    # distance, across clock changes too (checked), and it avoids a second guess at the zone.
    ev["event_end"] = (ev["event_start"] + (ev["end"] - ev["start"])).where(~day)
    ev["duration_published_h"] = ev["duration_pub_s"] / 3600.0
    elapsed = (ev["event_end"] - ev["event_start"]).dt.total_seconds() / 3600.0
    ev["duration_h"] = elapsed.where(~day, ev["duration_published_h"])
    bad = ev["duration_h"].notna() & ((ev["duration_h"] < 0) | (ev["duration_h"] >= 24 * 60))
    if bad.any():
        log.info("%s: dropped %d events shorter than 0 or longer than 60 days", source, int(bad.sum()))
    ev = ev[~bad]
    ev["site_name"] = ev["point"].astype(str).str.strip()
    ev["source"] = source
    log.info("%s: %d events, %d day records, %d overflows", source, len(ev), int(day.sum()), ev["site_id"].nunique())
    return ev[COLS + EXTRA].sort_values(["site_id", "event_start"]).reset_index(drop=True)


H = 3600


def _secs(s: pd.Series) -> np.ndarray:
    return (pd.to_datetime(s, utc=True) - pd.Timestamp("1970-01-01", tz="UTC")).dt.total_seconds().to_numpy(float)


def count_12_24(starts, ends) -> int:
    """Spills by the 12/24-hour block rule, as both workbooks' user guides state it (the EA's
    rule): the first discharge opens a 12-hour block that counts one spill; each following
    24-hour block with any discharge in it counts one more; a 24-hour block with none ends the
    run, and the next discharge opens a new 12-hour block."""
    s = _secs(pd.Series(starts))
    e = np.maximum(_secs(pd.Series(ends)), s)
    order = np.argsort(s, kind="stable")
    s, e = s[order], e[order]
    n, i = 0, 0
    while i < len(s):
        n += 1
        block_end, reach = s[i] + 12 * H, s[i]
        while True:
            while i < len(s) and s[i] < block_end:
                reach = max(reach, e[i])
                i += 1
            if reach > block_end or (i < len(s) and s[i] < block_end + 24 * H):
                n += 1
                block_end += 24 * H
            else:
                break
    return n


def yearly(ev: pd.DataFrame, rows: pd.DataFrame, source: str) -> pd.DataFrame:
    """One row per overflow (site_id) and sheet year. All measurement points of one Asset ID are
    pooled, as the live API has one row per Asset ID. Counts are per sheet year: a run of blocks
    does not carry over 1 January. Day records count as a discharge from midnight for their
    published duration; `day_records` says how many a year holds. Years an overflow is listed
    without events keep their status ("no events" counts as zero; "no data" is left empty)."""
    out = []
    for (sid, year), g in ev.groupby(["site_id", "year"]):
        ends = g["event_end"].fillna(g["event_start"] + pd.to_timedelta(g["duration_h"].fillna(0), unit="h"))
        out.append({"site_id": sid, "year": int(year), "spills_12_24": count_12_24(g["event_start"], ends),
                    "events": len(g), "hours": float(g["duration_h"].sum()),
                    "day_records": int(g["event_end"].isna().sum())})
    counted = pd.DataFrame(out, columns=["site_id", "year", "spills_12_24", "events", "hours", "day_records"])
    rank = rows.assign(r=rows["status"].map(STATUSES.index))
    meta = rank.groupby(["site_id", "year"]).agg(r=("r", "min"), days_data=("days_data", "max"),
                                                  points=("key", "nunique")).reset_index()
    meta["status"] = meta["r"].map(dict(enumerate(STATUSES)))
    df = meta.drop(columns="r").merge(counted, on=["site_id", "year"], how="left")
    df.loc[(df["status"] == "events") & df["events"].isna(), "status"] = "no events"   # all dropped as repeats
    zero = df["status"] == "no events"
    df.loc[zero, ["spills_12_24", "events", "day_records"]] = df.loc[zero, ["spills_12_24", "events", "day_records"]].fillna(0)
    df.loc[zero, "hours"] = df.loc[zero, "hours"].fillna(0.0)
    for c in ("spills_12_24", "events", "day_records"):
        df[c] = df[c].astype("Int64")
    df["source"] = source
    cols = ["site_id", "year", "source", "status", "spills_12_24", "events", "hours", "day_records", "days_data", "points"]
    return df[cols].sort_values(["site_id", "year"]).reset_index(drop=True)


def load(paths: dict[str, Path]) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    """(events, yearly, rows) for {source: workbook path}. `rows` keeps every listed row with its
    site_id, for the report."""
    read = {src: read_workbook(p) for src, p in paths.items()}
    evs, yrs, all_rows = [], [], []
    for src, (summary, rows) in read.items():
        others = [s for k, (s, _) in read.items() if k != src]
        rows = assign_sites(rows, [summary, *others])
        ev = events(rows, src)
        evs.append(ev)
        yrs.append(yearly(ev, rows, src))
        all_rows.append(rows.assign(source=src))
    return (pd.concat(evs, ignore_index=True), pd.concat(yrs, ignore_index=True),
            pd.concat(all_rows, ignore_index=True))


def report(ev: pd.DataFrame, yr: pd.DataFrame, rows: pd.DataFrame, api_ids: set[str] | None = None) -> pd.DataFrame:
    """Overflows and events by workbook and year: the numbers the S1 PR reports. Short events are
    counted on the published duration (computed where none is published)."""
    dur = ev["duration_published_h"].fillna(ev["duration_h"]) * 3600
    ev = ev.assign(short=dur.round() <= SHORT_S, day=ev["event_end"].isna())
    t = ev.groupby(["source", "year"]).agg(overflows=("site_id", "nunique"), events=("site_id", "size"),
                                           short_15m30s=("short", "sum"), day_records=("day", "sum"))
    t["listed"] = rows.groupby(["source", "year"])["site_id"].nunique()
    t["spills_12_24"] = yr.groupby(["source", "year"])["spills_12_24"].sum()
    t["hours"] = yr.groupby(["source", "year"])["hours"].sum().round(0)
    if api_ids is not None:
        g = ev.groupby(["source", "year"])
        t["in_api"] = g["site_id"].agg(lambda s: s[s.isin(api_ids)].nunique())
        t["published_in_api"] = g["asset_id_published"].agg(lambda s: s[s.isin(api_ids)].nunique())
        t["published_ids"] = g["asset_id_published"].nunique()
    return t.reset_index()
