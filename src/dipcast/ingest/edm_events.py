"""Pull event-level spill history (one row per discharge event).

United Utilities' files (config.EDM_EVENT_FEEDS) train the spill model. Hafren Dyfrdwy's 2025 file
(fetch_hd_events) is a Welsh test set only, kept apart in data/processed/hd_events.parquet so it
never enters training. Its credit: "Hafren Dyfrdwy Event Duration Monitoring 2025", Hafren
Dyfrdwy, published on Stream, licensed under CC BY 4.0
(https://creativecommons.org/licenses/by/4.0/).

Other companies' published histories (config.EDM_HISTORY_FEEDS: Northumbrian Water 2020-2024, Wessex
Water 2020-2023, South West Water January 2025 to May 2026) are read by fetch_history into their own
data/processed/stream_events_<company>.parquet (scripts/import_stream_events.py), never into
edm_events.parquet. Credits: HISTORY_CREDITS, all CC BY 4.0.

The Environment Agency's start/stop files for Dŵr Cymru Welsh Water's overflows in England
(read_ea_dcww_events) are a second Welsh-company test set, kept apart in
data/processed/dcww_ea_events.parquet. Their credit: "© Environment Agency copyright and/or
database right", Open Government Licence v3.0, from the EA dataset "Event Duration
Monitoring-Storm Overflow-Start/Stop Detailed Data".
"""

from __future__ import annotations

import io
import json
import logging
import re
import zipfile
from pathlib import Path
from urllib.parse import quote

import httpx
import numpy as np
import pandas as pd

from dipcast import config
from dipcast.arcgis import fetch_all
from dipcast.ingest.common import to_utc, write_parquet

log = logging.getLogger(__name__)

RENAME = {
    "SiteId": "site_id", "SiteName": "site_name", "OutfallLatitude": "lat",
    "OutfallLongitude": "lon", "EventStart": "event_start", "EventEnd": "event_end",
    "ReceivingWatercourse": "receiving_watercourse",
    # The same fields as other companies' history layers name them (HISTORY_FIELDS).
    "SiteID": "site_id", "Monitor_Unique_ID": "site_id", "Outfall_Latitude": "lat",
    "Outfall_Longitude": "lon", "Event_start": "event_start", "Event_end": "event_end",
    "ReceivingWaterCourse": "receiving_watercourse",
}
OUT_FIELDS = "SiteId,SiteName,OutfallLatitude,OutfallLongitude,EventStart,EventEnd,ReceivingWatercourse"
COLS = ["site_id", "site_name", "lat", "lon", "event_start", "event_end", "duration_h", "source"]


RAW_CACHE = config.CACHE / "edm"

# Hafren Dyfrdwy, the Severn Trent company that serves north-east and mid Wales. Not in Severn
# Trent's live feed or the EA's event files. Its 2025 file: 3,942 discharges at 43 site ids,
# EventStart/EventEnd as ISO text ending in "Z", latitude and longitude as text (checked 4 Oct 2026).
HD_SOURCE = "Hafren Dyfrdwy 2025"
HD_EVENTS_2025 = f"{config.STREAM_BASE}/HD_EDM_2025_Final_File/FeatureServer/0"
HD_LICENCE = "CC BY 4.0"
HD_CREDIT = ("Hafren Dyfrdwy Event Duration Monitoring 2025, Hafren Dyfrdwy, via Stream, "
             "licensed under CC BY 4.0 (https://creativecommons.org/licenses/by/4.0/)")


def _id_map() -> dict[str, str]:
    """Pre-2024 overflow IDs -> current IDs (Hub lookup table, then EA annual returns)."""
    m: dict[str, str] = {}
    ar_path = config.PROCESSED / "annual_returns.parquet"
    if ar_path.exists():
        ar = pd.read_parquet(ar_path, columns=["site_id", "old_site_id"]).dropna()
        m.update(dict(zip(ar["old_site_id"], ar["site_id"], strict=True)))
    lk_path = config.PROCESSED / "id_lookup.parquet"
    if lk_path.exists():
        lk = pd.read_parquet(lk_path).dropna()
        m.update(dict(zip(lk["old_site_id"], lk["site_id"], strict=True)))
    return m


def _normalise(df: pd.DataFrame, source: str) -> pd.DataFrame:
    df = df.rename(columns=RENAME)
    df["source"] = source
    # Hafren Dyfrdwy's layer stores the outfall position as text.
    df["lat"] = pd.to_numeric(df["lat"], errors="coerce")
    df["lon"] = pd.to_numeric(df["lon"], errors="coerce")
    m = _id_map()
    if m:
        mapped = df["site_id"].map(m)
        n = int(mapped.notna().sum())
        if n:
            log.info("%s: remapped %d rows from pre-2024 IDs", source, n)
        df["site_id"] = mapped.fillna(df["site_id"])
    # Each feed has one timestamp encoding; normalise before any concatenation.
    df["event_start"] = stream_times(df["event_start"])
    df["event_end"] = stream_times(df["event_end"])
    df = df.dropna(subset=["site_id", "event_start"])
    df["duration_h"] = (df["event_end"] - df["event_start"]).dt.total_seconds() / 3600.0
    df = df[(df["duration_h"].isna()) | ((df["duration_h"] >= 0) & (df["duration_h"] < 24 * 60))]
    # Northumbrian Water's layer has no site name.
    return df.reindex(columns=COLS)


def dmy_times(s: pd.Series) -> pd.Series:
    """UTC timestamps from "dd/mm/yyyy hh:mm" text, as Northumbrian Water's Stream layer gives its
    times. The text is GMT all year: every one of the EA's 106,346 Northumbrian starts for 2024 has a
    Stream start at the same overflow 0 or 1 minute earlier (the layer drops the seconds), in the
    summer-time months as in winter (checked 19 Oct 2026). Anything else is NaT."""
    t = pd.to_datetime(s, format="%d/%m/%Y %H:%M", errors="coerce")
    return t.dt.tz_localize("UTC").astype("datetime64[ns, UTC]")


def stream_times(s: pd.Series) -> pd.Series:
    """A Stream layer's time column as UTC: by dmy_times when most cells are day-first text, else to_utc."""
    if s.dtype.kind not in "iufM":
        text = s.dropna().astype(str)
        if len(text) and text.str.match(r"\d{2}/\d{2}/\d{4} \d{2}:\d{2}$").mean() > 0.5:
            return dmy_times(s)
    return to_utc(s)


def raw_cache_path(source: str) -> Path:
    """Where fetch_events keeps the rows of `source` as the layer gave them."""
    return RAW_CACHE / (source.replace(" ", "_") + ".parquet")


def fetch_events(feeds: dict[str, str] | None = None, refresh: bool = False,
                 out_fields: str = OUT_FIELDS) -> pd.DataFrame:
    feeds = feeds or config.EDM_EVENT_FEEDS
    RAW_CACHE.mkdir(parents=True, exist_ok=True)
    frames = []
    for source, url in feeds.items():
        raw_path = raw_cache_path(source)
        if raw_path.exists() and not refresh:
            raw = pd.read_parquet(raw_path)
            log.info("%s: %d events (cached)", source, len(raw))
        else:
            log.info("fetching %s ...", source)
            rows = fetch_all(
                url,
                geometry=False,
                out_fields=out_fields,
            )
            raw = pd.DataFrame(rows)
            raw.to_parquet(raw_path, index=False)
            log.info("%s: %d events", source, len(raw))
        frames.append(_normalise(raw, source))
    df = pd.concat(frames, ignore_index=True)
    return df.sort_values(["site_id", "event_start"]).reset_index(drop=True)


def fetch_hd_events(refresh: bool = False) -> pd.DataFrame:
    """Hafren Dyfrdwy's 2025 discharges in the same schema as fetch_events (CC BY 4.0; see the
    module docstring for the credit). For testing the model in Wales, never for training."""
    # The layer has no ReceivingWatercourse field, and ArcGIS refuses a query that names one.
    return fetch_events({HD_SOURCE: HD_EVENTS_2025}, refresh=refresh,
                        out_fields=OUT_FIELDS.removesuffix(",ReceivingWatercourse"))


# Other companies' event histories (config.EDM_HISTORY_FEEDS), never in edm_events.parquet.
# Each layer names its fields its own way; the object id is kept so a read can be checked for repeats.
# Northumbrian Water's layer has no site name or watercourse, and its times are text (dmy_times).
# South West Water's eventVerified is "TRUE" on every row and Wessex Water's EventVerified is empty on
# every row (checked 19 Oct 2026); both stay in the raw cache only.
HISTORY_FIELDS = {
    "Northumbrian Water 2020-2024": "ObjectId,Monitor_Unique_ID,Outfall_Latitude,Outfall_Longitude,Event_start,Event_end",
    "Wessex Water 2020-2023": "OBJECTID," + OUT_FIELDS + ",EventVerified",
    "South West Water 2025-2026": ("ObjectId,SiteID,SiteName,OutfallLatitude,OutfallLongitude,EventStart,EventEnd,"
                                   "ReceivingWaterCourse,eventVerified"),
}
HISTORY_COMPANY = {"Northumbrian Water 2020-2024": "Northumbrian Water", "Wessex Water 2020-2023": "Wessex Water",
                   "South West Water 2025-2026": "South West Water"}
HISTORY_CREDITS = {
    "Northumbrian Water 2020-2024": ("EDM Spill Data Northumbrian Water, Northumbrian Water, via Stream, "
                                     "licensed under CC BY 4.0 (https://creativecommons.org/licenses/by/4.0/)"),
    "Wessex Water 2020-2023": ("Wessex Water Event Duration Monitoring 2020-2023 © 2025 by Wessex Water, "
                               "licensed under CC BY 4.0 (https://creativecommons.org/licenses/by/4.0/)"),
    "South West Water 2025-2026": ("South West Water Storm Overflow EDM Monthly © 2025 by South West Water, via Stream, "
                                   "licensed under CC BY 4.0 (https://creativecommons.org/licenses/by/4.0/)"),
}


def fetch_history(source: str, refresh: bool = False) -> pd.DataFrame:
    """One of config.EDM_HISTORY_FEEDS in the fetch_events schema (CC BY 4.0; HISTORY_CREDITS)."""
    return fetch_events({source: config.EDM_HISTORY_FEEDS[source]}, refresh=refresh,
                        out_fields=HISTORY_FIELDS[source])


WESSEX_2024_SOURCE = "Wessex Water 2024"
WESSEX_2024_CREDIT = ("Wessex Water Event Duration Monitoring 2024 © 2025 by Wessex Water, "
                      "licensed under CC BY 4.0 (https://creativecommons.org/licenses/by/4.0/)")


def fetch_wessex_2024(refresh: bool = False) -> pd.DataFrame:
    """Wessex Water's own 2024 layer (config.WESSEX_EDM_2024), read only to build the id crosswalk."""
    return fetch_events({WESSEX_2024_SOURCE: config.WESSEX_EDM_2024}, refresh=refresh,
                        out_fields=HISTORY_FIELDS["Wessex Water 2020-2023"])


def crosswalk_by_events(a: pd.DataFrame, b: pd.DataFrame, min_share: float = 0.8) -> pd.DataFrame:
    """Ids of `a` -> ids of `b`, two records of the same events: each id of `a` goes to the id of `b`
    with the most events starting in the same minute as its own, if those are at least `min_share` of
    its events and of that id's. Both ways, because an id with one event can share its minute with a
    storm's start anywhere (Wessex 2024: 48 such pairs, a median 40 km apart, when only one way was
    asked). Columns: from_id, to_id, events, to_events, matched, share, to_share (one row per id of
    `a`; to_id is empty when either share is below min_share)."""
    def keyed(ev):
        return pd.DataFrame({"site_id": ev["site_id"].to_numpy(),
                             "minute": ev["event_start"].dt.floor("min").to_numpy()}).drop_duplicates()
    ka, kb = keyed(a), keyed(b)
    m = ka.merge(kb, on="minute", suffixes=("_a", "_b"))
    best = (m.groupby(["site_id_a", "site_id_b"]).size().rename("matched").reset_index()
            .sort_values("matched", ascending=False).drop_duplicates("site_id_a")
            .rename(columns={"site_id_a": "from_id", "site_id_b": "to_id"}))
    out = ka.groupby("site_id").size().rename("events").rename_axis("from_id").reset_index()
    out = out.merge(best, on="from_id", how="left")
    out["to_events"] = out["to_id"].map(kb.groupby("site_id").size())
    out["matched"] = out["matched"].fillna(0).astype(int)
    out["share"] = out["matched"] / out["events"]
    out["to_share"] = out["matched"] / out["to_events"]
    out.loc[(out["share"] < min_share) | ~(out["to_share"] >= min_share), "to_id"] = None
    return out.sort_values("from_id").reset_index(drop=True)


# The EA's start/stop files ("Event Duration Monitoring-Storm Overflow-Start/Stop Detailed Data"), one
# per company and calendar year, for the ten companies that operate in England. The DSP note
# ("Technical Detail DSP April 2026.docx", read 16 Oct 2026) gives one layout for all: column A the
# Unique ID (the annual returns' site_id), B the site name, C and D the activation's start and stop,
# headed "(GMT)"; one row per activation. The files differ in what holds that layout (checked
# 16 Oct 2026): some are CSV, some XLSX, and "Thames 2024 Detailed EDM Data.csv" is an XLSX workbook
# under a .csv name, so the reader looks at the bytes, not the extension. Times come as Excel dates,
# Excel serial numbers or text ("2024-01-01T20:14:32Z", "2024/01/01T20:14:32.000Z",
# "01/01/2024 20:14"). The files give no position; it comes from the annual returns.
# Dwr Cymru Welsh Water's files (overflows in England only) are a Welsh test set and never train.
# The EA publishes the files inside one zip per year ("2025 Detailed EDM Start-Stop Data.zip",
# "2026 year to August Detailed EDM Start-Stop Data.zip"); the per-company files are not published
# on their own any more. The year-to-date zip changes name as months are added, so its name is read
# from the dataset page's file list rather than written here.
EA_SOURCE = "EA"
EA_DATASET = "https://environment.data.gov.uk/dataset/e9677ac1-fd32-4ceb-88a0-2735be5f27c7"
EA_CREDIT = ("© Environment Agency copyright and/or database right. Licensed under the Open "
             "Government Licence v3.0 (https://www.nationalarchives.gov.uk/doc/open-government-licence/version/3/)")
EA_DCWW_SOURCE = "EA Welsh Water"
EA_DCWW_DATASET = EA_DATASET
EA_DCWW_CREDIT = EA_CREDIT
# File-name prefix -> the company name in annual_returns.parquet. "South West" is tried before
# "Southern" only for clarity: neither name starts with the other.
EA_COMPANIES = {
    "Anglian": "Anglian Water", "Northumbrian": "Northumbrian Water", "Severn Trent": "Severn Trent Water",
    "South West": "South West Water", "Southern": "Southern Water", "Thames": "Thames Water",
    "United Utilities": "United Utilities", "Wessex": "Wessex Water", "Yorkshire": "Yorkshire Water",
    "Welsh": "Dwr Cymru Welsh Water", "Dwr Cymru": "Dwr Cymru Welsh Water",
}
EA_TEST_ONLY = {"Dwr Cymru Welsh Water"}
# The platform's download route: a request that returns a short-lived signed link to the file.
EA_FILE_API = ("https://environment.data.gov.uk/file-management-open/data-sets/"
               "c95ecc47-fa70-4a24-b8cc-098232987526/files/{name}/download-url")
_EA_ZIP = re.compile(r"^(\d{4})\b.*Detailed EDM Start-Stop Data\.zip$", re.IGNORECASE)


def ea_zip_year(name: str) -> int | None:
    """The year of an EA yearly start/stop zip, from its name; None for any other file."""
    m = _EA_ZIP.match(Path(name).name)
    return int(m.group(1)) if m else None


def ea_zip_names() -> dict[int, str]:
    """Year -> name of each yearly start/stop zip in the dataset page's file list."""
    with httpx.Client(headers=config.EA_HEADERS, timeout=60, follow_redirects=True) as c:
        r = c.get(EA_DATASET)
        r.raise_for_status()
    m = re.search(r'<script id="__NEXT_DATA__"[^>]*>(.*?)</script>', r.text, re.S)
    if not m:
        raise RuntimeError(f"no file list on {EA_DATASET}; the page layout has changed")
    files = json.loads(m.group(1)).get("props", {}).get("pageProps", {}).get("files", [])
    out: dict[int, str] = {}
    for f in files:
        year = ea_zip_year(f.get("name", ""))
        if year is not None:
            if year in out:
                raise RuntimeError(f"two zips for {year}: {out[year]!r} and {f['name']!r}")
            out[year] = f["name"]
    return out


def download_ea_file(name: str, dest: Path) -> Path:
    """Download `name` from the EA dataset into the folder `dest` through the signed-link route."""
    if name != Path(name).name or "\\" in name:   # the name comes from a web page; keep it in `dest`
        raise ValueError(f"not a plain file name: {name!r}")
    with httpx.Client(headers=config.EA_HEADERS, timeout=600, follow_redirects=True) as c:
        r = c.get(EA_FILE_API.format(name=quote(name)))
        r.raise_for_status()
        f = c.get(r.json()["url"])
        f.raise_for_status()
    dest.mkdir(parents=True, exist_ok=True)
    (dest / name).write_bytes(f.content)
    log.info("downloaded %s (%d bytes)", name, len(f.content))
    return dest / name


def ea_yearly_zips(years, raw: Path, refresh: bool = False) -> dict[int, Path]:
    """Year -> local path of the EA's start/stop zip for each of `years`, kept in `raw`. A zip already
    there is used without asking the platform. Otherwise, or with `refresh`, the dataset's current
    list is read and its zips are downloaded (again, with `refresh`: the EA revises files in place
    under the same name). When a year has more than one
    local zip (a year-to-date zip renamed as months were added), the newest file is used."""
    years = sorted(set(years))
    local: dict[int, Path] = {}
    for p in sorted(raw.glob("*.zip"), key=lambda p: p.stat().st_mtime):
        y = ea_zip_year(p.name)
        if y in years:
            local[y] = p
    if refresh or set(local) != set(years):
        listed = ea_zip_names()
        missing = [y for y in years if y not in listed]
        if missing:
            raise RuntimeError(f"the EA dataset lists no start/stop zip for {missing}; it lists {listed}")
        for y in years:
            p = raw / listed[y]
            local[y] = p if p.exists() and not refresh else download_ea_file(listed[y], raw)
    return local


def ea_zip_members(zip_path) -> dict[str, list[tuple[Path, str]]]:
    """Company -> (zip, member) for every start/stop file inside a yearly zip. A data file whose
    name gives no known company stops the import: it would otherwise be dropped without notice."""
    zip_path = Path(zip_path)
    out: dict[str, list[tuple[Path, str]]] = {}
    with zipfile.ZipFile(zip_path) as zf:
        for info in zf.infolist():
            base = Path(info.filename).name
            if info.is_dir() or info.filename.startswith("__MACOSX/") or base.startswith((".", "~$")):
                continue
            if Path(base).suffix.lower() not in (".csv", ".xlsx"):
                continue
            company = ea_company(base)
            if company is None:
                raise ValueError(f"{zip_path.name}: no company for member {info.filename!r}")
            out.setdefault(company, []).append((zip_path, info.filename))
    return out


def ea_company(name: str) -> str | None:
    """The annual-returns company of an EA start/stop file, from its name."""
    for prefix, company in EA_COMPANIES.items():
        if name.startswith(prefix):
            return company
    return None


def _ea_name(src) -> str:
    """A start/stop file's own name, from a path or a (zip, member) pair."""
    return Path(src[1] if isinstance(src, tuple) else src).name


def _ea_bytes(src) -> bytes:
    """A start/stop file's bytes, from a path or a (zip, member) pair."""
    if isinstance(src, tuple):
        with zipfile.ZipFile(src[0]) as zf:
            return zf.read(src[1])
    return Path(src).read_bytes()


def _ea_offset(header) -> int:
    """Where the id column sits in a heading row: two columns before the first heading that names
    the start ("Discharge Start", or Thames's "Discharge Strat"). Some sheets put a permit number
    before the id; most begin with the id."""
    for i, cell in enumerate(header):
        s = str(cell or "").lower()
        if "start" in s or "strat" in s:
            return max(i - 2, 0)
    return 0


def _ea_rows(src) -> list[tuple]:
    """The id, name, start and stop cells of every row of every sheet of a start/stop file (XLSX by
    its bytes, else CSV), from a path or a (zip, member) pair."""
    raw = _ea_bytes(src)
    if raw[:4] == b"PK\x03\x04":
        from openpyxl import load_workbook  # dev dependency; only the off-CI import needs it
        # A file object, not the path: openpyxl refuses a workbook whose name ends ".csv".
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        rows = []
        for ws in wb.worksheets:
            off = None
            for r in ws.iter_rows(values_only=True):
                if off is None:
                    off = _ea_offset(r or ())
                if r and len(r) >= off + 4:
                    rows.append(tuple(r[off:off + 4]))
        wb.close()
        return rows
    for enc in ("utf-8-sig", "cp1252"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    head = pd.read_csv(io.StringIO(text), header=None, dtype=str, keep_default_na=False, nrows=1)
    off = _ea_offset(head.iloc[0].tolist()) if len(head) else 0
    df = pd.read_csv(io.StringIO(text), header=None, dtype=str, keep_default_na=False,
                     usecols=range(off, off + 4), skip_blank_lines=True)
    return list(df.itertuples(index=False, name=None))


_YMD = re.compile(r"^\d{4}[-/]\d{1,2}[-/]\d{1,2}")
_DMY = re.compile(r"^\d{1,2}[-/]\d{1,2}[-/]\d{4}")


def ea_times(values) -> pd.Series:
    """UTC timestamps from a column of mixed cells: datetime objects (no zone: read as GMT, as the
    files' headings say), Excel serial numbers, ISO text with "-" or "/" and an optional "Z", and
    day-first text ("01/01/2024 20:14"). Anything else is NaT. Cell by cell, because a column of
    Excel dates read whole would be taken for epoch numbers."""
    out = []
    for v in values:
        t = pd.NaT
        if v is None or (isinstance(v, float) and np.isnan(v)):
            pass
        elif hasattr(v, "isoformat"):
            t = pd.Timestamp(v)
        elif isinstance(v, (int, float, np.integer, np.floating)) and 20000 < float(v) < 80000:
            t = pd.Timestamp("1899-12-30") + pd.to_timedelta(float(v), unit="D")
        else:
            s = str(v).strip()
            try:
                if _YMD.match(s):
                    t = pd.Timestamp(s[:10].replace("/", "-") + s[10:])
                elif _DMY.match(s):
                    t = pd.to_datetime(s, dayfirst=True)
                elif s.replace(".", "", 1).isdigit() and 20000 < float(s) < 80000:
                    t = pd.Timestamp("1899-12-30") + pd.to_timedelta(float(s), unit="D")
            except (ValueError, OverflowError):
                t = pd.NaT
        if t is not pd.NaT and not pd.isna(t):
            t = t.tz_localize("UTC") if t.tzinfo is None else t.tz_convert("UTC")
            t = t.round("s")
        out.append(t)
    return pd.Series(pd.to_datetime(out, utc=True), dtype="datetime64[ns, UTC]")


def read_ea_events(paths, positions: pd.DataFrame | None = None, source: str = EA_SOURCE,
                   remap_ids: bool = True) -> pd.DataFrame:
    """The EA start/stop files `paths` (any company, CSV or XLSX, every sheet; each a path or a
    (zip, member) pair from ea_zip_members) in the fetch_events schema. Header rows and rows whose start is not a time are dropped; repeated rows (same id, start
    and stop) are kept once. Old ids are mapped to current ones (`_id_map`) when `remap_ids`.
    `positions` (site_id, lat, lon) defaults to the newest position of each id in
    annual_returns.parquet."""
    frames = []
    for src in paths:
        rows = [r for r in _ea_rows(src) if r[0] not in (None, "")]
        df = pd.DataFrame(rows, columns=["site_id", "site_name", "event_start", "event_end"])
        df["source"] = f"{source} ({_ea_name(src)})"
        frames.append(df)
    df = pd.concat(frames, ignore_index=True)
    df["site_id"] = df["site_id"].astype(str).str.strip()
    df["site_name"] = df["site_name"].astype(str).str.strip()
    df["event_start"] = ea_times(df["event_start"]).to_numpy()
    df["event_end"] = ea_times(df["event_end"]).to_numpy()
    df = df.dropna(subset=["event_start"])
    if remap_ids:
        m = _id_map()
        if m:
            df["site_id"] = df["site_id"].map(m).fillna(df["site_id"])
    df = df.drop_duplicates(["site_id", "event_start", "event_end"])
    if positions is None:
        ar = pd.read_parquet(config.PROCESSED / "annual_returns.parquet", columns=["site_id", "year", "lat", "lon"])
        positions = ar.dropna(subset=["site_id", "lat"]).sort_values("year").drop_duplicates("site_id", keep="last")
    df = df.merge(positions[["site_id", "lat", "lon"]], on="site_id", how="left")
    df["duration_h"] = (df["event_end"] - df["event_start"]).dt.total_seconds() / 3600.0
    df = df[(df["duration_h"].isna()) | ((df["duration_h"] >= 0) & (df["duration_h"] < 24 * 60))]
    return df[COLS].sort_values(["site_id", "event_start"]).reset_index(drop=True)


def read_ea_dcww_events(paths, positions: pd.DataFrame | None = None) -> pd.DataFrame:
    """The EA's Welsh Water workbooks (`paths`) in the fetch_events schema (read_ea_events). For
    testing the model, never for training."""
    return read_ea_events(paths, positions=positions, source=EA_DCWW_SOURCE, remap_ids=False)


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(name)s: %(message)s")
    ev = fetch_events()
    write_parquet(ev, config.PROCESSED / "edm_events.parquet")
    print(ev.groupby("source").agg(events=("site_id", "size"), sites=("site_id", "nunique"),
                                    first=("event_start", "min"), last=("event_start", "max")))
